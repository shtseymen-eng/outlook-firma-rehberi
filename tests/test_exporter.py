from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook
import pytest

from outlook_rehberi.exporter import (
    build_workbook,
    default_export_name,
    save_workbook_atomic,
)
from outlook_rehberi.models import (
    Category,
    Contact,
    ScanResult,
    ScanStats,
    ScanWarning,
    SourceRef,
)


NOW = datetime(2026, 10, 24, 14, 5, tzinfo=timezone.utc)
SOURCE = SourceRef("a1", "owner@example.com", "inbox", ("Gelen", "Müşteriler"))


def contact(
    email: str,
    company: str,
    category: Category,
    *,
    name: str = "",
) -> Contact:
    return Contact(
        email=email,
        domain=email.rsplit("@", 1)[1],
        company=company,
        category=category,
        name=name,
        sources=(SOURCE,),
    )


def scan_result(*, complete: bool = True) -> ScanResult:
    return ScanResult(
        contacts=(),
        stats=ScanStats(
            started_at=NOW,
            completed_at=NOW,
            accounts=2,
            folders=7,
            messages=125,
            unique_contacts=7,
        ),
        complete=complete,
        warnings=(
            ScanWarning("blocked@example.com", ("Arşiv",), "İzin reddedildi"),
        )
        if not complete
        else (),
    )


def sample_contacts() -> list[Contact]:
    return [
        contact("z@polisan.com.tr", "POLİSAN", Category.POLISAN_POLIPORT),
        contact("a@polisan.com.tr", "POLİSAN", Category.POLISAN_POLIPORT),
        contact("port@poliport.com", "POLİPORT", Category.POLISAN_POLIPORT),
        contact("customs@acme-ygm.com", "ACME YGM", Category.CUSTOMS_YGM),
        contact("two@vendor.com", "VENDOR LTD.", Category.CORPORATE),
        contact("one@vendor.com", "VENDOR LTD.", Category.CORPORATE),
        contact("person@gmail.com", "GMAIL", Category.PERSONAL),
    ]


def test_workbook_has_exact_sheet_order_headers_and_summary():
    workbook = build_workbook(sample_contacts(), scan_result())

    assert workbook.sheetnames == [
        "Özet",
        "Polisan-Poliport",
        "Gümrük-YGM",
        "Diğer Kurumsal",
        "Kişisel Mailler",
    ]
    details = workbook["Diğer Kurumsal"]
    assert [details.cell(3, column).value for column in range(1, 10)] == [
        "Ad Soyad",
        "Görev / Unvan",
        "Domain",
        "Firma",
        "E-posta",
        "Cep Telefonu",
        "Sabit Telefon",
        "Kaynak Hesap(lar)",
        "Kaynak Klasör(ler)",
    ]
    summary_values = {
        cell.value
        for row in workbook["Özet"].iter_rows()
        for cell in row
        if cell.value is not None
    }
    assert "Toplam Tekil E-posta" in summary_values
    assert 7 in summary_values


def test_workbook_sorts_email_rows_and_keeps_polisan_poliport_as_two_blocks():
    workbook = build_workbook(sample_contacts(), scan_result())
    sheet = workbook["Polisan-Poliport"]

    assert [sheet[f"E{row}"].value for row in range(4, 7)] == [
        "port@poliport.com",
        "a@polisan.com.tr",
        "z@polisan.com.tr",
    ]
    assert sheet["D4"].value == "POLİPORT"
    assert sheet["D5"].value == "POLİSAN"
    assert "D5:D6" in {str(item) for item in sheet.merged_cells.ranges}
    assert "D4:D6" not in {str(item) for item in sheet.merged_cells.ranges}


def test_incomplete_scan_is_marked_in_summary():
    workbook = build_workbook(sample_contacts(), scan_result(complete=False))
    values = [
        str(cell.value)
        for row in workbook["Özet"].iter_rows()
        for cell in row
        if cell.value is not None
    ]

    assert any("EKSİK" in value for value in values)
    assert any("blocked@example.com" in value for value in values)


def test_default_export_name_contains_local_timestamp():
    assert default_export_name(NOW) == "Outlook_Firma_Rehberi_2026-10-24_14-05.xlsx"


def test_workbook_is_globally_unique_grouped_merged_and_atomic(tmp_path):
    contacts = sample_contacts()
    contacts.append(contacts[-1])
    workbook = build_workbook(contacts, scan_result())
    destination = tmp_path / "rehber.xlsx"

    assert save_workbook_atomic(workbook, destination) == destination
    loaded = load_workbook(destination)
    detail_names = loaded.sheetnames[1:]
    emails = [
        sheet.cell(row, 5).value
        for name in detail_names
        for sheet in [loaded[name]]
        for row in range(4, sheet.max_row + 1)
        if sheet.cell(row, 5).value
    ]
    assert len(emails) == len(set(emails)) == 7
    assert "D4:D5" in {
        str(item) for item in loaded["Diğer Kurumsal"].merged_cells.ranges
    }

    original_bytes = destination.read_bytes()

    class BrokenWorkbook:
        def save(self, path):
            Path(path).write_bytes(b"partial")
            raise OSError("disk full")

    with pytest.raises(OSError, match="disk full"):
        save_workbook_atomic(BrokenWorkbook(), destination)

    assert destination.read_bytes() == original_bytes
    assert list(tmp_path.glob("*.tmp.xlsx")) == []

