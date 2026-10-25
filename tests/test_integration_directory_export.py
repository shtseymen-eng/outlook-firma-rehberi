from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

from outlook_rehberi.classifier import normalize_email
from outlook_rehberi.exporter import SHEET_ORDER, build_workbook, save_workbook_atomic
from outlook_rehberi.models import (
    AccountRef,
    CancellationToken,
    FolderRef,
    GraphMessage,
    GraphMessagePage,
    Participant,
    ScanSelection,
)
from outlook_rehberi.scanner import ScanCallbacks, ScanService
from outlook_rehberi.storage import LocalStore


NOW = datetime(2026, 10, 24, 12, 0, tzinfo=timezone.utc)


class FixtureGraph:
    def __init__(self, pages) -> None:
        self.pages = pages
        self.calls: list[tuple[str, str]] = []

    def iter_messages(
        self,
        account: AccountRef,
        folder: FolderRef,
        cancellation: CancellationToken,
        delta_link: str | None = None,
    ):
        del delta_link
        self.calls.append((account.account_id, folder.folder_id))
        for page in self.pages[(account.account_id, folder.folder_id)]:
            cancellation.raise_if_cancelled()
            yield page


def account(account_id: str, email: str) -> AccountRef:
    return AccountRef(account_id, email, email, "tenant")


def folder(
    owner: AccountRef,
    folder_id: str,
    name: str,
    parent_id: str | None,
    path: tuple[str, ...],
) -> FolderRef:
    return FolderRef(owner.account_id, folder_id, name, parent_id, path, 0, False)


def incoming(message_id: str, sender: str, body: str = "") -> GraphMessage:
    return GraphMessage(
        message_id=message_id,
        sender=Participant(sender, sender.split("@", 1)[0]),
        body_text=body,
        received_at=NOW,
    )


def sent(
    message_id: str,
    owner: AccountRef,
    *recipients: str,
) -> GraphMessage:
    return GraphMessage(
        message_id=message_id,
        sender=Participant(owner.email, owner.display_name),
        to=tuple(Participant(value, value.split("@", 1)[0]) for value in recipients),
        sent_at=NOW,
        is_sent=True,
    )


def detail_emails(workbook) -> dict[str, list[str]]:
    result: dict[str, list[str]] = {}
    for sheet_name in SHEET_ORDER[1:]:
        sheet = workbook[sheet_name]
        result[sheet_name] = [
            str(sheet.cell(row, 5).value)
            for row in range(4, sheet.max_row + 1)
            if sheet.cell(row, 5).value
        ]
    return result


def test_two_account_nested_mailbox_scan_exports_complete_categorized_workbook(
    tmp_path: Path,
) -> None:
    first = account("a1", "birinci@bizim.com")
    second = account("a2", "ikinci@bizim.com")
    inbox = folder(first, "inbox", "Gelen", None, ("Gelen",))
    deleted = folder(first, "deleted", "Silinen", None, ("Silinen",))
    deep = folder(
        first,
        "deep",
        "2026",
        "deleted",
        ("Silinen", "Arşiv", "Müşteriler", "2026"),
    )
    sent_folder = folder(second, "sent", "Gönderilen", None, ("Gönderilen",))

    pages = {
        ("a1", "inbox"): (
            GraphMessagePage(
                (
                    incoming("1", "ali@polisan.com.tr", "Saygılarımla\nAli Veli\nPOLİSAN A.Ş."),
                    incoming("2", "ayse@polisan.com.tr", "Ayşe Kaya\nPOLİSAN A.Ş."),
                    incoming(
                        "3",
                        "operasyon@abc.com.tr",
                        "İyi çalışmalar\nABC Gümrük Müşavirliği A.Ş.\nOperasyon Uzmanı",
                    ),
                    incoming("4", "kisi@gmail.com"),
                ),
                next_link="inbox-page-2",
            ),
            GraphMessagePage(
                (
                    incoming("5", "ortak@firma.com", "Ortak Kişi\nFİRMA A.Ş."),
                    incoming(
                        "6",
                        "normal@normal.com",
                        "Merhaba\nOn Friday eski@example.com wrote:\nESKİ GÜMRÜK A.Ş.",
                    ),
                    incoming("7", "bozuk-adres", "Geçersiz kayıt"),
                ),
                delta_link="delta-inbox",
            ),
        ),
        ("a1", "deleted"): (
            GraphMessagePage(
                (incoming("8", "liman@poliport.com", "POLİPORT A.Ş."),),
                delta_link="delta-deleted",
            ),
        ),
        ("a1", "deep"): (
            GraphMessagePage(
                (incoming("9", "ortak@firma.com", "Güncel İmza\nFİRMA A.Ş."),),
                delta_link="delta-deep",
            ),
        ),
        ("a2", "sent"): (
            GraphMessagePage(
                (
                    sent(
                        "10",
                        second,
                        "ortak@firma.com",
                        "ozel@yandex.com",
                        first.email,
                        "takma@bizim.com",
                    ),
                ),
                delta_link="delta-sent",
            ),
        ),
    }
    graph = FixtureGraph(pages)
    store = LocalStore(tmp_path / "rehber.sqlite3")
    store.set_aliases("a1", {"takma@bizim.com"})
    scanner = ScanService(graph, store)

    result = scanner.run(
        (
            ScanSelection(first, (inbox, deleted, deep), frozenset()),
            ScanSelection(second, (sent_folder,), frozenset()),
        ),
        full_rebuild=True,
        cancellation=CancellationToken(),
        callbacks=ScanCallbacks(),
    )

    assert result.complete is True
    assert result.stats.accounts == 2
    assert result.stats.folders == 4
    assert result.stats.messages == 10
    assert len(result.contacts) == 8
    shared = next(item for item in result.contacts if item.email == "ortak@firma.com")
    assert {source.account_id for source in shared.sources} == {"a1", "a2"}
    assert all(item.email not in {first.email, second.email, "takma@bizim.com"} for item in result.contacts)

    destination = save_workbook_atomic(
        build_workbook(result.contacts, result),
        tmp_path / "iki-hesap-rehberi.xlsx",
    )
    workbook = load_workbook(destination)

    assert tuple(workbook.sheetnames) == SHEET_ORDER
    assert workbook["Özet"]["A2"].value == "TAMAMLANDI"
    emails = detail_emails(workbook)
    assert set(emails["Polisan-Poliport"]) == {
        "ali@polisan.com.tr",
        "ayse@polisan.com.tr",
        "liman@poliport.com",
    }
    assert emails["Gümrük-YGM"] == ["operasyon@abc.com.tr"]
    assert set(emails["Kişisel Mailler"]) == {"kisi@gmail.com", "ozel@yandex.com"}
    assert set(emails["Diğer Kurumsal"]) == {"normal@normal.com", "ortak@firma.com"}
    flattened = [normalize_email(value) for values in emails.values() for value in values]
    assert len(flattened) == len(set(flattened)) == 8
    assert "D5:D6" in {str(item) for item in workbook["Polisan-Poliport"].merged_cells.ranges}
    assert destination.is_file()
