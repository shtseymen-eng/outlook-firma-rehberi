from __future__ import annotations

from collections import Counter, defaultdict
from datetime import datetime
import os
from pathlib import Path
import tempfile
from typing import Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .classifier import classify_contact, company_group_key, normalize_email, normalize_search_text
from .models import Category, Contact, ScanResult


SHEET_ORDER = (
    "Özet",
    "Polisan-Poliport",
    "Gümrük-YGM",
    "Diğer Kurumsal",
    "Kişisel Mailler",
)
CATEGORY_SHEETS = {
    Category.POLISAN_POLIPORT: "Polisan-Poliport",
    Category.CUSTOMS_YGM: "Gümrük-YGM",
    Category.CORPORATE: "Diğer Kurumsal",
    Category.PERSONAL: "Kişisel Mailler",
}
HEADERS = (
    "Ad Soyad",
    "Görev / Unvan",
    "Domain",
    "Firma",
    "E-posta",
    "Cep Telefonu",
    "Sabit Telefon",
    "Kaynak Hesap(lar)",
    "Kaynak Klasör(ler)",
)
WIDTHS = (24, 28, 24, 28, 36, 18, 18, 32, 42)
ACCENTS = {
    "Polisan-Poliport": "1D4ED8",
    "Gümrük-YGM": "B45309",
    "Diğer Kurumsal": "0F766E",
    "Kişisel Mailler": "7C3AED",
}


def default_export_name(now: datetime) -> str:
    return f"Outlook_Firma_Rehberi_{now.strftime('%Y-%m-%d_%H-%M')}.xlsx"


def _unique_contacts(contacts: Sequence[Contact]) -> tuple[Contact, ...]:
    unique: dict[str, Contact] = {}
    for contact in contacts:
        email = normalize_email(contact.email)
        if email is not None:
            unique.setdefault(email, contact)
    return tuple(unique[email] for email in sorted(unique))


def _border() -> Border:
    side = Side(style="thin", color="CBD5E1")
    return Border(left=side, right=side, top=side, bottom=side)


def _style_cell(cell, *, fill: str = "FFFFFF", bold: bool = False, color: str = "0F172A") -> None:
    cell.font = Font(name="Calibri", size=10, bold=bold, color=color)
    cell.fill = PatternFill("solid", fgColor=fill)
    cell.alignment = Alignment(vertical="center", wrap_text=True)
    cell.border = _border()


def _accounts(contact: Contact) -> str:
    return " / ".join(sorted({source.account_email for source in contact.sources}))


def _folders(contact: Contact) -> str:
    return " / ".join(
        sorted({" › ".join(source.folder_path) for source in contact.sources})
    )


def _company_label(contacts: Sequence[Contact]) -> str:
    labels = [contact.company.strip() for contact in contacts if contact.company.strip()]
    if not labels:
        return contacts[0].domain.split(".")[0].upper()
    counts = Counter(normalize_search_text(label) for label in labels)
    best_key = min(counts, key=lambda key: (-counts[key], key))
    return min(label for label in labels if normalize_search_text(label) == best_key)


def _write_detail_sheet(workbook: Workbook, title: str, contacts: Sequence[Contact]) -> None:
    sheet = workbook.create_sheet(title)
    accent = ACCENTS[title]
    sheet.sheet_view.showGridLines = False
    sheet.merge_cells("A1:I1")
    sheet["A1"] = f"OUTLOOK FİRMA REHBERİ — {title.upper()}"
    _style_cell(sheet["A1"], fill=accent, bold=True, color="FFFFFF")
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 30
    sheet.merge_cells("A2:I2")
    sheet["A2"] = f"Toplam tekil e-posta: {len(contacts)}"
    _style_cell(sheet["A2"], fill="E2E8F0", color="334155")
    sheet["A2"].alignment = Alignment(horizontal="center", vertical="center")
    for column, (header, width) in enumerate(zip(HEADERS, WIDTHS), start=1):
        cell = sheet.cell(3, column, header)
        _style_cell(cell, fill=accent, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.row_dimensions[3].height = 24

    groups: dict[str, list[Contact]] = defaultdict(list)
    for contact in contacts:
        groups[company_group_key(contact)[0]].append(contact)

    row = 4
    for group_index, domain in enumerate(sorted(groups)):
        people = sorted(groups[domain], key=lambda item: item.email)
        start_row = row
        fill = "EFF6FF" if group_index % 2 == 0 else "FFFFFF"
        label = _company_label(people)
        for contact in people:
            values = (
                contact.name,
                contact.title,
                contact.domain,
                label,
                contact.email,
                " / ".join(contact.mobile_numbers),
                " / ".join(contact.landline_numbers),
                _accounts(contact),
                _folders(contact),
            )
            for column, value in enumerate(values, start=1):
                cell = sheet.cell(row, column)
                cell.value = value
                if isinstance(value, str):
                    # Contact data originates in untrusted message metadata/signatures.
                    # Force text so formula-looking values cannot execute in Excel.
                    cell.data_type = "s"
                _style_cell(cell, fill=fill)
            row += 1
        end_row = row - 1
        if end_row > start_row:
            sheet.merge_cells(start_row=start_row, start_column=4, end_row=end_row, end_column=4)
            sheet.merge_cells(start_row=start_row, start_column=3, end_row=end_row, end_column=3)
            sheet.cell(start_row, 3).alignment = Alignment(vertical="center", wrap_text=True)
            sheet.cell(start_row, 4).alignment = Alignment(vertical="center", wrap_text=True)
    sheet.freeze_panes = "A4"
    sheet.auto_filter.ref = f"A3:I{max(3, row - 1)}"


def _write_summary(workbook: Workbook, contacts: Sequence[Contact], result: ScanResult) -> None:
    sheet = workbook.active
    sheet.title = "Özet"
    sheet.sheet_view.showGridLines = False
    sheet.merge_cells("A1:I1")
    sheet["A1"] = "OUTLOOK FİRMA REHBERİ — ÖZET"
    _style_cell(sheet["A1"], fill="1E3A5F", bold=True, color="FFFFFF")
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 34
    status = "TAMAMLANDI" if result.complete else "EKSİK TARAMA — yeniden deneyin"
    sheet.merge_cells("A2:I2")
    sheet["A2"] = status
    _style_cell(
        sheet["A2"],
        fill="DCFCE7" if result.complete else "FEE2E2",
        bold=True,
        color="166534" if result.complete else "991B1B",
    )
    counts = Counter(contact.category or classify_contact(contact) for contact in contacts)
    rows = (
        ("Toplam Tekil E-posta", len(contacts)),
        ("Polisan-Poliport", counts[Category.POLISAN_POLIPORT]),
        ("Gümrük-YGM", counts[Category.CUSTOMS_YGM]),
        ("Diğer Kurumsal", counts[Category.CORPORATE]),
        ("Kişisel Mailler", counts[Category.PERSONAL]),
        ("Taranan Hesap", result.stats.accounts),
        ("Taranan Klasör", result.stats.folders),
        ("Okunan Mesaj", result.stats.messages),
    )
    for row, (label, value) in enumerate(rows, start=4):
        sheet.cell(row, 1, label)
        sheet.cell(row, 2, value)
        _style_cell(sheet.cell(row, 1), fill="F1F5F9", bold=True)
        _style_cell(sheet.cell(row, 2), fill="FFFFFF")
    warning_row = 4 + len(rows) + 2
    for warning in result.warnings:
        sheet.merge_cells(start_row=warning_row, start_column=1, end_row=warning_row, end_column=9)
        sheet.cell(
            warning_row,
            1,
            f"{warning.account_email} — {' › '.join(warning.folder_path)} — {warning.message}",
        )
        _style_cell(sheet.cell(warning_row, 1), fill="FEF3C7", color="92400E")
        warning_row += 1
    sheet.column_dimensions["A"].width = 28
    sheet.column_dimensions["B"].width = 18


def build_workbook(contacts: Sequence[Contact], result: ScanResult) -> Workbook:
    unique = _unique_contacts(contacts)
    categorized: dict[Category, list[Contact]] = defaultdict(list)
    for contact in unique:
        category = contact.category or classify_contact(contact)
        categorized[category].append(contact)

    workbook = Workbook()
    _write_summary(workbook, unique, result)
    for category in (
        Category.POLISAN_POLIPORT,
        Category.CUSTOMS_YGM,
        Category.CORPORATE,
        Category.PERSONAL,
    ):
        _write_detail_sheet(workbook, CATEGORY_SHEETS[category], categorized[category])
    return workbook


def save_workbook_atomic(workbook: Workbook, destination: Path) -> Path:
    destination = Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(
        dir=destination.parent,
        prefix=f".{destination.stem}.",
        suffix=".tmp.xlsx",
    )
    os.close(descriptor)
    temporary = Path(temporary_name)
    try:
        workbook.save(temporary)
        validation = load_workbook(temporary, read_only=True)
        validation.close()
        os.replace(temporary, destination)
    except Exception:
        temporary.unlink(missing_ok=True)
        raise
    return destination

